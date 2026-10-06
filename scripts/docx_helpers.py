"""Helpers for editing Word CVs built on a table template, with python-docx.

The template assumed here: a profile of plain paragraphs, four capability blocks
each written as one paragraph whose first run is a bold label and whose second run
is the body, and a career table whose cells hold the role entries as paragraphs styled
"List Paragraph". Adjust the style name if your template differs.
"""
import copy
from docx import Document
from docx.text.paragraph import Paragraph

BULLET_STYLE = "List Paragraph"


def set_text(paragraph, text):
    """Replace a paragraph's text, keeping the formatting of its first run."""
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run._r.getparent().remove(run._r)


def set_block(paragraph, label, body):
    """Rewrite a capability block: bold label, then a body of lowercase things."""
    runs = paragraph.runs
    runs[0].text = label
    runs[0].bold = True
    runs[1].text = "  " + body
    runs[1].bold = False
    for run in runs[2:]:
        run._r.getparent().remove(run._r)


def delete(paragraph):
    paragraph._p.getparent().remove(paragraph._p)


def cells(container):
    """Yield every table cell in a document, including nested tables."""
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield cell
                yield from cells(cell)


def bullets(cell):
    return [p for p in cell.paragraphs if p.style.name == BULLET_STYLE and p.text.strip()]


W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"


def _strip_copied_ids(element):
    """A copied paragraph must not carry the original's bookmarks or paragraph ids; Word reports duplicates as corruption."""
    for tag in ("bookmarkStart", "bookmarkEnd"):
        for el in element.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + tag):
            el.getparent().remove(el)
    for el in [element] + list(element.iter()):
        for attr in ("paraId", "textId"):
            el.attrib.pop("{%s}%s" % (W14_NS, attr), None)


def set_bullets(cell, texts):
    """Rewrite a cell's bullet list to exactly these texts, one theme per bullet."""
    existing = bullets(cell)
    while len(existing) < len(texts):
        new = copy.deepcopy(existing[-1]._p)
        _strip_copied_ids(new)
        existing[-1]._p.addnext(new)
        existing.append(Paragraph(new, existing[-1]._parent))
    for paragraph, text in zip(existing, texts):
        set_text(paragraph, text)
    for paragraph in existing[len(texts):]:
        delete(paragraph)


def find_paragraph(document, prefix):
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise KeyError(prefix)


if __name__ == "__main__":
    import sys
    if len(sys.argv) != 2:
        sys.exit("usage: python docx_helpers.py path/to/cv.docx  (prints the document's paragraphs)")
    doc = Document(sys.argv[1])
    for p in doc.paragraphs:
        if p.text.strip():
            print(p.text[:120])
