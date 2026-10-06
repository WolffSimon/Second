"""Add a tools line under named education entries of a table-based CV.

Usage: python add_education_tools.py path/to/cv.docx path/to/education_tools.json

The JSON maps the start of an institution's line, as it appears in the CV, to the line to add beneath it:

    {"Example University": "Stack: Python, pandas, scikit-learn."}

The script is idempotent: if the row beneath an institution already starts with "Stack:", it is replaced, not duplicated.
"""
import copy
import json
import sys

from docx import Document
from docx.table import Table, _Row


def set_text(paragraph, text):
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run._r.getparent().remove(run._r)
    for run in paragraph.runs:
        run.bold = False
        run.italic = True


def last_cell_text(row):
    return row.cells[-1].text.strip()


def insert_or_replace(inner, row_idx, text):
    rows = list(inner.rows)
    if row_idx + 1 < len(rows) and last_cell_text(rows[row_idx + 1]).startswith("Stack:"):
        target = rows[row_idx + 1].cells[-1].paragraphs[0]
        set_text(target, text)
        return "replaced"
    new = copy.deepcopy(rows[row_idx]._tr)
    rows[row_idx]._tr.addnext(new)
    new_row = _Row(new, inner)
    cells = new_row.cells
    paragraphs = cells[-1].paragraphs
    set_text(paragraphs[0], text)
    for p in paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for c in cells[:-1]:
        for p in c.paragraphs:
            for r in p.runs:
                r.text = ""
    return "added"


def nested_tables(document):
    for element in document.element.body.iterchildren():
        if not element.tag.endswith("}tbl"):
            continue
        for row in Table(element, document).rows:
            for cell in row.cells:
                yield from cell.tables


def main(cv_path, config_path):
    mapping = json.load(open(config_path, encoding="utf-8"))
    document = Document(cv_path)
    done = {}
    for inner in nested_tables(document):
        # Work bottom-up so inserted rows do not shift indexes still to be visited.
        for idx in range(len(inner.rows) - 1, -1, -1):
            text = last_cell_text(inner.rows[idx])
            for prefix, line in mapping.items():
                if text.startswith(prefix) and prefix not in done:
                    done[prefix] = insert_or_replace(inner, idx, line)
    document.save(cv_path)
    for prefix in mapping:
        print(f"{prefix}: {done.get(prefix, 'not found')}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
