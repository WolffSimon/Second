"""Export a table-template CV to plain text a reviewer can read: profile, then each role with title, company, dates and entries.

Usage: python export_cv_text.py path/to/cv.docx > cv.txt
"""
import sys
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

if len(sys.argv) != 2:
    sys.exit(__doc__)
d = Document(sys.argv[1])
out = []

def nested_header(cell):
    lines = []
    for nt in cell.tables:
        for r in nt.rows:
            for c in r.cells:
                t = c.text.strip()
                if t and t not in lines:
                    lines.append(t)
    return lines

for child in d.element.body.iterchildren():
    tag = child.tag.split('}')[1]
    if tag == 'p':
        p = Paragraph(child, d)
        t = p.text.strip()
        if not t:
            continue
        if p.style.name.startswith('Heading'):
            out.append("\n## " + t)
        elif p.style.name == 'List Paragraph':
            out.append("- " + t)
        else:
            out.append(t + "\n")
    elif tag == 'tbl':
        tbl = Table(child, d)
        for row in tbl.rows:
            cells = row.cells
            if len(cells) >= 2 and cells[1].tables:
                dates = cells[0].text.strip()
                hdr = nested_header(cells[1])
                out.append("\n### " + " | ".join(hdr) + (" | " + dates if dates else ""))
                for p in cells[1].paragraphs:
                    if p.text.strip():
                        out.append("- " + p.text.strip())
            else:
                for c in cells:
                    for p in c.paragraphs:
                        if p.text.strip() and p.style.name.startswith('Heading 1'):
                            out.append("# " + p.text.strip())
                        elif p.text.strip():
                            out.append(p.text.strip())
                    for nt in c.tables:
                        for r in nt.rows:
                            line = " ".join(x.text.strip() for x in r.cells if x.text.strip())
                            if line:
                                out.append(line)
print("\n".join(out))
