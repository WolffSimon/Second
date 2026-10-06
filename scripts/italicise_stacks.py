"""Italicise every paragraph beginning 'Stack:' in a CV, in body and in tables."""
import sys
from docx import Document
def cells(c):
    for t in c.tables:
        for r in t.rows:
            for cell in r.cells:
                yield cell
                yield from cells(cell)
def main(path):
    d=Document(path); n=0
    paras=list(d.paragraphs)+[p for c in cells(d) for p in c.paragraphs]
    for p in paras:
        if p.text.strip().startswith("Stack:"):
            for r in p.runs: r.italic=True
            n+=1
    d.save(path); print(f"{n} stack lines italicised")
if __name__=="__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: python italicise_stacks.py path/to/cv.docx")
    main(sys.argv[1])
